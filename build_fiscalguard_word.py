from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"C:\FACULDADE\fabrica\FiscalGuard_Tech_Trabalho_Fabrica_Projetos_VI_final_contabilidade.docx")
DOWNLOADS = Path(r"C:\Users\rayss\Downloads")

REPLACEMENTS = {
    "Indice": "Índice",
    "Definicao": "Definição",
    "definicao": "definição",
    "Solucao": "Solução",
    "solucao": "solução",
    "Analise": "Análise",
    "analise": "análise",
    "Requisitos Nao Funcionais": "Requisitos Não Funcionais",
    "Requisito nao funcional": "Requisito não funcional",
    "nao": "não",
    "Nao": "Não",
    "Negocio": "Negócio",
    "negocio": "negócio",
    "Sequencia": "Sequência",
    "sequencia": "sequência",
    "Atividades": "Atividades",
    "Consideracoes": "Considerações",
    "Fabrica": "Fábrica",
    "Aceleradora": "Aceleradora",
    "Cristovam": "Cristóvam",
    "composicao": "composição",
    "optantes": "optantes",
    "situacao": "situação",
    "Situacao": "Situação",
    "cadastral": "cadastral",
    "opcao": "opção",
    "Opcao": "Opção",
    "Simples Nacional": "Simples Nacional",
    "pendencias": "pendências",
    "Pendencias": "Pendências",
    "relatorios": "relatórios",
    "Relatorio": "Relatório",
    "relatorio": "relatório",
    "recomendacoes": "recomendações",
    "informacoes": "informações",
    "integracoes": "integrações",
    "Integracoes": "Integrações",
    "Receita Federal": "Receita Federal",
    "municipais": "municipais",
    "notificacoes": "notificações",
    "varredura": "varredura",
    "frequencia": "frequência",
    "Frequency": "Frequency",
    "seguranca": "segurança",
    "Seguranca": "Segurança",
    "auditabilidade": "auditabilidade",
    "Auditabilidade": "Auditabilidade",
    "manutenibilidade": "manutenibilidade",
    "Manutenibilidade": "Manutenibilidade",
    "confiabilidade": "confiabilidade",
    "Confiabilidade": "Confiabilidade",
    "autenticacao": "autenticação",
    "decisao": "decisão",
    "decisoes": "decisões",
    "gestao": "gestão",
    "Gestao": "Gestão",
    "inteligencia": "inteligência",
    "Inteligencia": "Inteligência",
    "assíncrono": "assíncrono",
    "historico": "histórico",
    "Historico": "Histórico",
    "historicos": "históricos",
    "tecnica": "técnica",
    "tecnico": "técnico",
    "tributaria": "tributária",
    "tributarias": "tributárias",
    "tributario": "tributário",
    "notificacao": "notificação",
    "acao": "ação",
    "Acao": "Ação",
    "acoes": "ações",
    "pratico": "prático",
    "publico": "público",
    "Publico": "Público",
    "usuario": "usuário",
    "usuarios": "usuários",
    "Empresario": "Empresário",
    "empresario": "empresário",
    "empresarios": "empresários",
    "escritorios": "escritórios",
    "contabeis": "contábeis",
    "Contador": "Contador",
    "indisponibilidade": "indisponibilidade",
    "CNPJs": "CNPJs",
    "Irregularidade": "Irregularidade",
    "irregularidade": "irregularidade",
    "disponiveis": "disponíveis",
    "disponivel": "disponível",
    "servicos": "serviços",
    "Servico": "Serviço",
    "servico": "serviço",
    "Descricao": "Descrição",
    "descricao": "descrição",
    "Codigo": "Código",
    "criterio": "critério",
    "Criterio": "Critério",
    "criterios": "critérios",
    "criterio": "critério",
    "saude": "saúde",
    "excessiva": "excessiva",
    "visao": "visão",
    "Visao": "Visão",
    "prevencao": "prevenção",
    "Prevencao": "Prevenção",
    "previsibilidade": "previsibilidade",
    "propria": "própria",
    "modulos": "módulos",
    "Modulo": "Módulo",
    "dominio": "domínio",
    "operacoes": "operações",
    "associacoes": "associações",
    "relacao": "relação",
    "relacoes": "relações",
    "Parametros": "Parâmetros",
    "parametros": "parâmetros",
    "Assinatura": "Assinatura",
    "Assinaturas": "Assinaturas",
    "comunicacao": "comunicação",
    "composicao": "composição",
    "Composicao": "Composição",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=105, bottom=70, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None, header_fill="1F4E79"):
    table.style = "Table Grid"
    table.autofit = False
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if widths and c < len(widths):
                cell.width = widths[c]
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.name = "Aptos"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(8.7)
            if r == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif r % 2 == 0:
                set_cell_shading(cell, "F4F8FB")


def set_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in [("Title", 26), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6 if style_name != "Title" else 14)
        style.paragraph_format.keep_with_next = True


def set_section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def set_section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "FiscalGuard Tech | Trabalho de Fabrica de Projetos VI"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(80, 80, 80)


def add_picture_center(doc, path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=width)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Descricao"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table, [Cm(4.2), Cm(11.8)])
    doc.add_paragraph()


def add_matrix(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(table, widths)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_manual_sumario(doc):
    doc.add_heading("Indice", level=1)
    items = [
        "1 Definicao do Software",
        "2 Definicao do Problema",
        "3 Definicao da Solucao",
        "4 Analise de Requisitos",
        "5 Requisitos Funcionais",
        "6 Requisitos Nao Funcionais",
        "7 Regras de Negocio",
        "8 Diagrama de Caso de Uso",
        "9 Diagrama de Classes",
        "10 Diagrama de Sequencia",
        "11 Diagramas de Atividades",
        "12 Consideracoes Finais",
    ]
    for item in items:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(2)


def fix_text(text):
    for old, new in sorted(REPLACEMENTS.items(), key=lambda kv: len(kv[0]), reverse=True):
        text = text.replace(old, new)
    return text


def apply_language_polish(doc):
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text:
                run.text = fix_text(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text:
                            run.text = fix_text(run.text)
    for section in doc.sections:
        for p in section.footer.paragraphs:
            for run in p.runs:
                if run.text:
                    run.text = fix_text(run.text)


def build():
    doc = Document()
    set_styles(doc)
    set_section_portrait(doc.sections[0])
    add_footer(doc.sections[0])

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph("TRABALHO DE FABRICA DE PROJETOS VI")
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph("LEAVENING - Aceleradora de Startups")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.runs[0].font.size = Pt(13)
    p3 = doc.add_paragraph("Prof. Cristovam")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.runs[0].font.size = Pt(12)
    for _ in range(4):
        doc.add_paragraph()
    p4 = doc.add_paragraph("FiscalGuard Tech")
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.runs[0].bold = True
    p4.runs[0].font.size = Pt(22)
    p5 = doc.add_paragraph("Plataforma SaaS de monitoramento fiscal preventivo para escritórios de contabilidade")
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.runs[0].font.size = Pt(12)
    for _ in range(6):
        doc.add_paragraph()
    p6 = doc.add_paragraph("Grupo FiscalGuard Tech")
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.runs[0].bold = True
    p7 = doc.add_paragraph("Integrantes: preencher conforme a composicao oficial da equipe")
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_manual_sumario(doc)
    doc.add_page_break()

    doc.add_heading("1 Definicao do Software", level=1)
    add_paragraph(doc, "O FiscalGuard Tech e uma plataforma SaaS voltada ao monitoramento fiscal preventivo para escritorios de contabilidade que atendem empresas optantes pelo Simples Nacional. O sistema acompanha a situacao cadastral e fiscal dos CNPJs da carteira, identifica pendencias, calcula riscos, emite alertas antecipados e apresenta recomendacoes em linguagem acessivel para a equipe contabil e seus clientes.")
    add_paragraph(doc, "A proposta central e transformar a gestao fiscal dos escritorios de uma postura reativa para uma rotina preventiva. Em vez de o contador descobrir irregularidades apenas apos notificacoes, exclusao do regime tributario ou cobrancas acumuladas, a plataforma organiza sinais de risco por cliente e orienta a tomada de decisao antes que o problema gere impacto financeiro.")
    add_kv_table(doc, [
        ("Nome do software", "FiscalGuard Tech"),
        ("Categoria", "Plataforma SaaS de inteligencia fiscal preventiva"),
        ("Publico alvo", "Escritorios de contabilidade que monitoram carteiras de PMEs optantes pelo Simples Nacional."),
        ("Atores principais", "Contador parceiro, equipe do escritorio contabil, empresario cliente e administrador FiscalGuard."),
        ("Integracoes previstas", "Bases publicas da Receita Federal para dados cadastrais, fontes municipais de ISS quando houver disponibilidade tecnica e gateway de e-mail ou SMS para notificacoes."),
        ("Modelo de receita", "Assinatura mensal para escritorios de contabilidade, com planos escalonados por faixa de CNPJs monitorados."),
    ])

    doc.add_heading("2 Definicao do Problema", level=1)
    add_paragraph(doc, "Embora o Simples Nacional seja divulgado como regime simplificado, os escritorios de contabilidade precisam acompanhar prazos, debitos, inconsistencias cadastrais, comunicacoes fiscais e exigencias que podem variar conforme municipio, atividade economica e situacao do CNPJ de cada cliente. Em carteiras com muitos CNPJs, a prevencao fiscal pode se tornar manual, dispersa e dificil de priorizar.")
    add_paragraph(doc, "Esse atraso aumenta a chance de multas, juros, perda de previsibilidade financeira, dificuldade de emitir documentos fiscais, risco de exclusao do Simples Nacional e sobrecarga da equipe contabil. O problema central do escritorio e monitorar varios clientes ao mesmo tempo, com informacoes espalhadas em diferentes fontes e necessidade de resposta rapida.")
    add_bullets(doc, [
        "Baixa visibilidade sobre pendencias fiscais dos clientes antes da notificacao formal.",
        "Dificuldade de priorizar quais CNPJs da carteira exigem acao imediata.",
        "Dependencia de verificacoes manuais em portais e bases publicas.",
        "Ausencia de historico consolidado para avaliar evolucao de riscos por cliente.",
        "Custo potencial elevado quando o escritorio ou o cliente age somente depois do problema instalado.",
    ])

    doc.add_heading("3 Definicao da Solucao", level=1)
    add_paragraph(doc, "A solucao proposta combina monitoramento continuo de carteiras, classificacao de pendencias, alertas e recomendacoes corretivas. O escritorio cadastra ou importa os CNPJs dos clientes, o sistema valida os dados, consulta fontes externas, consolida informacoes e apresenta um painel com a situacao fiscal de cada empresa. Quando identifica irregularidades ou sinais de risco, a plataforma registra a pendencia, gera alerta e sugere encaminhamentos.")
    add_paragraph(doc, "O diferencial do FiscalGuard Tech esta no foco preventivo para escritorios contabeis que atendem empresas do Simples Nacional. A plataforma nao pretende substituir a responsabilidade tecnica da contabilidade, mas apoiar a equipe com informacoes organizadas, historico de consultas, relatorios de risco e notificacoes objetivas para orientar seus clientes.")
    add_matrix(doc, ["Modulo", "Finalidade", "Resultado esperado"], [
        ("Monitoramento e gestao", "Consultar situacao cadastral, validar opcao pelo Simples Nacional, identificar pendencias e gerar relatorios por cliente.", "Visao atualizada da carteira fiscal e priorizacao dos CNPJs com maior risco."),
        ("Inteligencia e prevencao", "Enviar alertas, sugerir acoes corretivas, simular multas e realizar analise retroativa.", "Reducao de surpresa fiscal e apoio a decisoes antecipadas."),
        ("Administracao", "Gerenciar planos por faixa de CNPJs e configurar parametros de varredura.", "Operacao sustentavel da plataforma e controle dos limites de alerta do escritorio."),
    ], [Cm(4.0), Cm(7.0), Cm(5.0)])

    doc.add_heading("4 Analise de Requisitos", level=1)
    add_paragraph(doc, "A analise de requisitos foi organizada a partir dos atores do sistema e dos principais fluxos de uso. O contador parceiro utiliza o painel para consultar empresas da carteira, acompanhar status, receber alertas e orientar os clientes. O empresario cliente aparece como beneficiario das informacoes e pode receber comunicacoes sobre pendencias. O administrador FiscalGuard configura parametros, planos e assinaturas.")
    add_matrix(doc, ["Ator", "Responsabilidade no sistema", "Necessidade atendida"], [
        ("Contador parceiro", "Cadastrar, importar e acompanhar CNPJs de clientes, gerar relatorios e consultar recomendacoes.", "Priorizar clientes e prevenir problemas antes de vencimentos ou notificacoes."),
        ("Equipe do escritorio", "Visualizar painel da carteira, tratar pendencias e registrar acompanhamentos.", "Distribuir a rotina de monitoramento fiscal entre responsaveis internos."),
        ("Empresario cliente", "Receber comunicacoes e orientacoes encaminhadas pelo escritorio.", "Entender riscos sem depender de linguagem tecnica excessiva."),
        ("Administrador FG", "Gerenciar assinaturas por faixa de CNPJs, parametrizar varreduras e manter limites de alerta.", "Controlar a operacao comercial e tecnica da plataforma."),
        ("Receita Federal e Prefeitura", "Disponibilizar dados consultados pelas integracoes.", "Fornecer sinais de regularidade, pendencia e situacao cadastral."),
        ("Gateway de e-mail SMS", "Enviar notificacoes ao usuario.", "Garantir comunicacao rapida sobre riscos e falhas de envio."),
    ], [Cm(3.3), Cm(7.4), Cm(5.3)])

    doc.add_heading("5 Requisitos Funcionais", level=1)
    add_matrix(doc, ["Codigo", "Requisito funcional", "Prioridade"], [
        ("RF01", "Permitir consulta da situacao cadastral de empresas da carteira a partir dos CNPJs informados.", "Alta"),
        ("RF02", "Validar o formato do CNPJ antes de acionar consultas externas.", "Alta"),
        ("RF03", "Consultar dados fiscais em servicos da Receita Federal e da prefeitura, quando disponiveis.", "Alta"),
        ("RF04", "Verificar se a empresa atende aos criterios de opcao pelo Simples Nacional.", "Alta"),
        ("RF05", "Identificar debitos, inconsistencias e pendencias tributarias.", "Alta"),
        ("RF06", "Registrar pendencias com descricao, valor, status e data de vencimento.", "Alta"),
        ("RF07", "Gerar alerta de irregularidade, prazo ou risco conforme parametros definidos.", "Alta"),
        ("RF08", "Enviar notificacoes por e-mail ou SMS e registrar sucesso ou falha de envio.", "Media"),
        ("RF09", "Gerar relatorio de risco com data de geracao e score calculado.", "Media"),
        ("RF10", "Sugerir acoes corretivas conforme o tipo de problema identificado.", "Media"),
        ("RF11", "Simular multas e juros associados a uma pendencia.", "Media"),
        ("RF12", "Permitir analise retroativa para demonstrar riscos historicos ao usuario.", "Media"),
        ("RF13", "Permitir ao administrador gerenciar assinaturas e planos por faixa de CNPJs monitorados.", "Alta"),
        ("RF14", "Permitir configuracao da frequencia de varredura e limites de alerta.", "Alta"),
    ], [Cm(2.0), Cm(11.5), Cm(2.5)])

    doc.add_heading("6 Requisitos Nao Funcionais", level=1)
    add_matrix(doc, ["Codigo", "Requisito nao funcional", "Criterio de qualidade"], [
        ("RNF01", "Usabilidade", "A interface deve usar linguagem acessivel e orientar usuarios sem conhecimento tecnico tributario."),
        ("RNF02", "Seguranca", "Dados de usuarios, empresas e consultas devem ser protegidos por autenticacao e controle de acesso por perfil."),
        ("RNF03", "Disponibilidade", "O sistema deve manter acesso estavel ao painel e tratar indisponibilidade de APIs externas com mensagens claras."),
        ("RNF04", "Escalabilidade", "A arquitetura deve suportar crescimento por quantidade de CNPJs monitorados e por escritorios contabeis."),
        ("RNF05", "Auditabilidade", "Consultas, alertas e envios devem ficar registrados para rastreabilidade."),
        ("RNF06", "Manutenibilidade", "Integracoes externas devem ficar isoladas em servicos especificos para facilitar ajustes quando APIs mudarem."),
        ("RNF07", "Desempenho", "Consultas e paineis devem responder em tempo adequado para uso recorrente, com processamento assíncrono para varreduras."),
        ("RNF08", "Confiabilidade", "Falhas de notificacao e indisponibilidade de fonte externa devem ser registradas e apresentadas ao usuario."),
    ], [Cm(2.0), Cm(4.2), Cm(9.8)])

    doc.add_page_break()
    doc.add_heading("7 Regras de Negocio", level=1)
    add_matrix(doc, ["Codigo", "Regra de negocio"], [
        ("RN01", "Somente CNPJs com formato valido devem seguir para consulta em fontes externas."),
        ("RN02", "Uma empresa pode estar nos status REGULAR, IRREGULAR ou EM_ANALISE conforme os dados consolidados."),
        ("RN03", "Pendencias devem iniciar com status ABERTA e somente passar para RESOLVIDA apos confirmacao ou nova verificacao."),
        ("RN04", "Alertas podem ser classificados como IRREGULARIDADE, PRAZO ou RISCO."),
        ("RN05", "O envio de alerta deve registrar tanto o sucesso quanto a falha de notificacao."),
        ("RN06", "O score de risco deve considerar historico fiscal, pendencias abertas, valores estimados e vencimentos proximos."),
        ("RN07", "A analise retroativa deve usar historico fiscal para demonstrar eventos passados e oportunidades de prevencao."),
        ("RN08", "Assinaturas devem controlar plano, valor, faixa de CNPJs monitorados, datas de inicio e fim e situacao ativa."),
        ("RN09", "Administradores podem configurar frequencia de varredura e limite de alertas; usuarios comuns nao podem alterar esses parametros globais."),
        ("RN10", "A plataforma deve indicar recomendacoes como apoio a tomada de decisao, preservando a responsabilidade tecnica do contador quando aplicavel."),
    ], [Cm(2.0), Cm(14.0)])

    # Landscape diagrams
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(sec)
    add_footer(sec)
    doc.add_heading("8 Diagrama de Caso de Uso", level=1)
    add_paragraph(doc, "O diagrama de caso de uso apresenta os atores externos, os principais modulos e as relacoes entre funcionalidades de monitoramento, inteligencia preventiva e administracao.")
    add_picture_center(doc, DOWNLOADS / "diagrama_caso_uso_fiscalguard.png", Inches(10.2))
    add_caption(doc, "Figura 1 - Diagrama de casos de uso detalhado do FiscalGuard Tech.")

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(sec)
    add_footer(sec)
    doc.add_heading("9 Diagrama de Classes", level=1)
    add_paragraph(doc, "O diagrama de classes corrigido representa as principais entidades do dominio, seus atributos, operacoes e associacoes. A classe Empresa concentra o relacionamento com monitoramento, pendencias, alertas, relatorios, assinatura e validacao do Simples Nacional.")
    add_picture_center(doc, DOWNLOADS / "diagrama_classe_fiscalguard_new.png", Inches(10.4))
    add_caption(doc, "Figura 2 - Diagrama de classes corrigido do FiscalGuard Tech.")

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(sec)
    add_footer(sec)
    doc.add_heading("10 Diagrama de Sequencia", level=1)
    add_paragraph(doc, "O diagrama de sequencia demonstra o fluxo de monitoramento fiscal preventivo: o usuario acessa o painel, o sistema consulta dados externos, registra pendencia quando ha irregularidade, gera alerta e envia notificacao ao usuario.")
    add_picture_center(doc, DOWNLOADS / "fiscalguard_diagrama_Sequencia.png", Inches(10.0))
    add_caption(doc, "Figura 3 - Diagrama de sequencia do monitoramento fiscal preventivo.")

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(sec)
    add_footer(sec)
    doc.add_heading("11 Diagramas de Atividades", level=1)
    add_paragraph(doc, "Os diagramas de atividades detalham os fluxos internos das principais funcionalidades. Eles evidenciam pontos de decisao, caminhos alternativos, registros de falha e etapas de exibicao de resultados ao usuario.")

    activities = [
        ("Consultar Situacao Cadastral", "cnpj.png"),
        ("Identificar Pendencias", "pendencias.png"),
        ("Validar Simples Nacional", "validar_simples_nacional.png"),
        ("Enviar Alerta", "alertas.png"),
        ("Gerar Relatorio de Risco", "relatorios.png"),
        ("Sugerir Acao Corretiva", "acao_corretiva.png"),
        ("Simular Multas e Juros", "simular_multas.png"),
        ("Analise Retroativa", "analise_retroativa.png"),
        ("Gerenciar Assinatura", "gerenciar_assinaturas.png"),
        ("Configurar Parametros", "parametros.png"),
    ]
    for i in range(0, len(activities), 2):
        table = doc.add_table(rows=2, cols=2)
        table.autofit = False
        for c in range(2):
            if i + c >= len(activities):
                continue
            title, filename = activities[i + c]
            title_cell = table.rows[0].cells[c]
            title_cell.text = title
            title_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_cell.paragraphs[0].runs[0].bold = True
            img_cell = table.rows[1].cells[c]
            img_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_cell.paragraphs[0].add_run().add_picture(str(DOWNLOADS / "fiscalguard_diagrama_atividade" / filename), width=Inches(2.55))
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=60, start=60, bottom=60, end=60)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                # white borders for layout grid
                set_cell_borders(cell, color="FFFFFF", size="2")
        doc.add_paragraph()

    doc.add_heading("12 Consideracoes Finais", level=1)
    add_paragraph(doc, "O FiscalGuard Tech apresenta uma proposta coerente para um problema especifico e recorrente: a dificuldade dos escritorios de contabilidade em acompanhar riscos fiscais de varios clientes de forma preventiva. A combinacao de monitoramento, historico, alertas e recomendacoes torna o produto util para a operacao contabil e melhora a comunicacao com as empresas atendidas.")
    add_paragraph(doc, "A continuidade do projeto deve priorizar a validacao das fontes de dados, a definicao de planos por faixa de CNPJs monitorados, a analise de responsabilidade sobre alertas preventivos e a comparacao com funcionalidades ja presentes em softwares contabeis. Esses pontos fortalecem a viabilidade tecnica e comercial da solucao.")

    doc.core_properties.title = "FiscalGuard Tech Trabalho de Fabrica de Projetos VI"
    doc.core_properties.subject = "Documento profissional do projeto FiscalGuard Tech"
    doc.core_properties.author = "FiscalGuard Tech"
    apply_language_polish(doc)
    doc.save(OUT)


if __name__ == "__main__":
    build()
